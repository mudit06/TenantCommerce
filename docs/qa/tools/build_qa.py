"""Builds docs/qa from the case files: one Markdown regression script per module, the coverage
page and the Excel workbook. Run from the repo root:  python docs/qa/tools/build_qa.py
Needs Python 3.10+ and openpyxl (pip install openpyxl). Results of runs live in results.json."""

import json
import sys
from datetime import date
from pathlib import Path

sys.path.insert(0, str(Path(__file__).parent))

from openpyxl import Workbook  # noqa: E402
from openpyxl.formatting.rule import CellIsRule, FormulaRule  # noqa: E402
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side  # noqa: E402
from openpyxl.utils import get_column_letter  # noqa: E402
from openpyxl.worksheet.datavalidation import DataValidation  # noqa: E402

import cases_e2e  # noqa: E402
import cases_sa  # noqa: E402
import cases_vc  # noqa: E402
from qa_model import ADMIN, PORTALS, PRIORITIES, STATUSES, STORE, TYPES, check  # noqa: E402

ROOT = Path(__file__).resolve().parents[1]           # docs/qa
REPO = ROOT.parents[1]
MODULES = cases_sa.MODULES + cases_vc.MODULES + cases_e2e.MODULES
CASES = cases_sa.CASES + cases_vc.CASES + cases_e2e.CASES
check(CASES, MODULES)
BY_ID = {m.id: m for m in MODULES}
RESULTS = json.loads((ROOT / 'tools' / 'results.json').read_text(encoding='utf-8')) \
    if (ROOT / 'tools' / 'results.json').exists() else {}
DEFECTS = RESULTS.pop('_defects', [])
FOLDER = {'Super Admin': 'super-admin', 'Vendor CMS': 'vendor-cms', 'End-to-End': 'e2e'}


def script_path(module) -> str:
    return f'docs/qa/regression/{FOLDER[module.portal]}/{module.id}-{module.slug}.md'


def module_label(module) -> str:
    return f'{module.id} · {module.name}'


def result(case) -> dict:
    return RESULTS.get(case.id, {})


# ------------------------------------------------------------------ Markdown scripts
RUN_WITH_CHROME = f'''## How to run this script with Claude in Chrome

1. Start the QA app: in Claude Code start the `app-qa` launch entry (or `DATABASE_URI=…/tenantecom_qa pnpm dev:fresh`).
   Wait until `{STORE}` answers (the first page compiles for up to a minute).
2. Open Chrome with the Claude extension and say: "Run the regression script
   `<this file>` against localhost. Follow each step, verify the expected result on screen
   before moving on, and record Actual result and Status for each test case."
3. Claude signs in only with the test accounts in `docs/qa/README.md` (Test accounts). It never
   uses real passwords, keys or card numbers, and asks before anything that sends a message
   outside localhost.
4. For every step, wait for a visible condition (a heading, a message, a row) instead of a fixed
   delay. After every save, reload and check the value stayed. If a page shows an error or the
   expected text never appears, mark the case **Fail** (or **Blocked** when a precondition is
   missing), note what was seen, and continue with the next case.
5. Copy each Actual result and Status into the workbook (`docs/qa/TenantEcom-test-cases.xlsx`,
   same Test Case ID), or paste the run's output into the shared sheet.
'''


def write_scripts() -> None:
    for portal, folder in FOLDER.items():
        (ROOT / 'regression' / folder).mkdir(parents=True, exist_ok=True)
    for module in MODULES:
        cases = [c for c in CASES if c.module == module.id]
        lines = [
            f'# {module.id} {module.name}',
            '',
            f'Regression script, generated from `docs/qa/tools` (edit the case files there, then run '
            f'`python docs/qa/tools/build_qa.py`).',
            '',
            '| | |',
            '|---|---|',
            f'| Portal | {module.portal} |',
            f'| Purpose | {module.purpose} |',
            f'| Runs as | {module.role} |',
            f'| Start at | `{module.route}` |',
            f'| Environment | Local QA: admin `{ADMIN}`, store `{STORE}`, database `tenantecom_qa` |',
            f'| Spec | {", ".join(f"`{s}`" for s in module.screens) or "—"} (docs/screens) |',
            f'| Depends on | {", ".join(module.depends_on) or "—"} |',
            f'| Test cases | {len(cases)} ({sum(1 for c in cases if c.smoke)} in the smoke run) |',
            '',
            '## Before you start',
            '',
            *[f'- {p}' for p in module.preconditions],
            '',
            RUN_WITH_CHROME,
            '## Test cases',
            '',
        ]
        for c in cases:
            r = result(c)
            lines += [
                f'### {c.id} · {c.sub}: {c.func}',
                '',
                '| Field | Value |',
                '|---|---|',
                f'| Module | {module_label(module)} |',
                f'| Test case ID | {c.id} |',
                f'| Scenario and objective | {c.scenario} |',
                f'| Preconditions | {c.pre} |',
                f'| Test data | {c.data} |',
                f'| Role | {c.role or module.role} |',
                f'| Priority | {c.priority} |',
                f'| Type | {c.type} |',
                '',
                '**Steps**',
                '',
                *[f'{i}. {s}' for i, s in enumerate(c.steps, 1)],
                f'{len(c.steps) + 1}. Verify the expected result below before going on.',
                '',
                f'**Expected result:** {c.expected}',
                '',
                f'**Actual result:** {r.get("actual", "_(fill in when run)_")}',
                '',
                f'**Status:** {r.get("status", "Not Run")}'
                + (f' ({r.get("by")}, {r.get("date")})' if r.get('by') else ''),
                '',
                f'**Execution notes / defect:** {"; ".join(x for x in [c.notes, r.get("notes", ""), r.get("defect", "")] if x) or "—"}',
                '',
            ]
        (REPO / script_path(module)).write_text('\n'.join(lines), encoding='utf-8')


# ------------------------------------------------------------------ Coverage page
def write_coverage() -> None:
    lines = [
        '# Module inventory and coverage matrix',
        '',
        'Generated by `docs/qa/tools/build_qa.py`. Every module of both portals, the screens it covers,',
        'its regression script and how many manual test cases cover each kind of check.',
        '',
    ]
    types_short = ['Functional', 'Negative', 'Validation', 'Boundary', 'Integration',
                   'Security/Permissions', 'Data persistence', 'UI/UX']
    for portal in PORTALS:
        mods = [m for m in MODULES if m.portal == portal]
        lines += [f'## {portal}', '',
                  '| Module | Screens | Cases | Critical | Smoke | ' + ' | '.join(types_short) + ' | Script |',
                  '|---|---|---|---|---|' + '---|' * len(types_short) + '---|']
        for m in mods:
            cs = [c for c in CASES if c.module == m.id]
            counts = [sum(1 for c in cs if t in c.type.split(', ')) for t in types_short]
            lines.append(
                f'| {m.id} {m.name} | {", ".join(m.screens)} | {len(cs)} | '
                f'{sum(1 for c in cs if c.priority == "Critical")} | {sum(1 for c in cs if c.smoke)} | '
                + ' | '.join(str(n) if n else '·' for n in counts)
                + f' | [{m.id}](../../{script_path(m)}) |'.replace('../../docs/qa/', '')
            )
        total = [c for c in CASES if BY_ID[c.module].portal == portal]
        lines += ['', f'{len(mods)} modules, {len(total)} test cases.', '']
    lines += [
        '## Dependencies between modules',
        '',
        '| Module | Needs first |',
        '|---|---|',
        *[f'| {m.id} {m.name} | {", ".join(m.depends_on) or "—"} |' for m in MODULES if m.depends_on],
        '',
    ]
    (ROOT / 'coverage.md').write_text('\n'.join(lines), encoding='utf-8')


# ------------------------------------------------------------------ Workbook
HEAD = PatternFill('solid', fgColor='1F2937')
HEAD_FONT = Font(bold=True, color='FFFFFF')
THIN = Side(style='thin', color='D1D5DB')
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(wrap_text=True, vertical='top')
COLS = [
    ('Test Case ID', 13), ('Application/Portal', 13), ('Module Name', 26), ('Submodule/Feature', 18),
    ('Functionality', 22), ('Test Scenario', 38), ('Preconditions', 30), ('Test Data', 28),
    ('Step-by-Step Instructions', 55), ('Expected Result', 48), ('Actual Result', 36), ('Status', 12),
    ('Priority', 10), ('Test Type', 18), ('Executed By', 16), ('Execution Date', 13),
    ('Defect/Bug ID', 12), ('Remarks/Additional Notes', 34), ('Role', 18), ('Smoke Run', 9),
    ('Regression Script', 40),
]
STATUS_COL = 'L'


def header(ws, cols) -> None:
    for i, (name, width) in enumerate(cols, 1):
        cell = ws.cell(row=1, column=i, value=name)
        cell.fill, cell.font, cell.border = HEAD, HEAD_FONT, BOX
        cell.alignment = Alignment(wrap_text=True, vertical='center')
        ws.column_dimensions[get_column_letter(i)].width = width
    ws.row_dimensions[1].height = 32
    ws.freeze_panes = 'B2'


def status_colours(ws, col: str, last: int) -> None:
    rng = f'{col}2:{col}{max(last, 2)}'
    for value, colour in [('Pass', 'C6EFCE'), ('Fail', 'FFC7CE'), ('Blocked', 'FFEB9C'),
                          ('Not Run', 'E5E7EB'), ('Not Applicable', 'DBEAFE')]:
        ws.conditional_formatting.add(
            rng, CellIsRule(operator='equal', formula=[f'"{value}"'], fill=PatternFill('solid', fgColor=colour)))


def list_validation(ws, values, rng) -> None:
    dv = DataValidation(type='list', formula1='"' + ','.join(values) + '"', allow_blank=True)
    dv.error, dv.errorTitle = 'Pick a value from the list', 'Invalid value'
    ws.add_data_validation(dv)
    dv.add(rng)


def case_sheet(wb, title, portal) -> int:
    ws = wb.create_sheet(title)
    header(ws, COLS)
    rows = [c for c in CASES if BY_ID[c.module].portal == portal]
    for r, c in enumerate(rows, 2):
        m = BY_ID[c.module]
        res = result(c)
        types = c.type.split(', ')
        steps = '\n'.join(f'{i}. {s}' for i, s in enumerate(c.steps, 1))
        steps += f'\n{len(c.steps) + 1}. Verify the expected result.'
        remarks = '; '.join(x for x in [
            c.notes, ('Also: ' + ', '.join(types[1:])) if len(types) > 1 else '', res.get('notes', '')] if x)
        values = [c.id, portal if portal != 'End-to-End' else 'End-to-End', module_label(m), c.sub, c.func,
                  c.scenario, c.pre, c.data, steps, c.expected, res.get('actual', ''),
                  res.get('status', 'Not Run'), c.priority, types[0], res.get('by', ''),
                  res.get('date', ''), res.get('defect', ''), remarks, c.role or m.role,
                  'Yes' if c.smoke else 'No', script_path(m)]
        for col, v in enumerate(values, 1):
            cell = ws.cell(row=r, column=col, value=v)
            cell.alignment, cell.border = WRAP, BOX
    last = len(rows) + 1
    ws.auto_filter.ref = f'A1:{get_column_letter(len(COLS))}{last}'
    list_validation(ws, STATUSES, f'L2:L{last + 200}')
    list_validation(ws, PRIORITIES, f'M2:M{last + 200}')
    list_validation(ws, TYPES, f'N2:N{last + 200}')
    list_validation(ws, PORTALS, f'B2:B{last + 200}')
    list_validation(ws, ['Yes', 'No'], f'T2:T{last + 200}')
    status_colours(ws, STATUS_COL, last + 200)
    for col, colour in [('M', None)]:
        ws.conditional_formatting.add(f'M2:M{last + 200}', CellIsRule(
            operator='equal', formula=['"Critical"'], font=Font(bold=True, color='B91C1C')))
    return last


def summary_sheet(wb) -> None:
    ws = wb.active
    ws.title = 'Test Summary'
    ws['A1'] = 'TenantEcom: manual and regression test summary'
    ws['A1'].font = Font(bold=True, size=14)
    ws['A2'] = f'Generated {date.today().isoformat()} from docs/qa/tools. Figures recalculate as Status cells change.'
    ws['A2'].font = Font(italic=True, color='6B7280')
    sheets = [('Super Admin', "'Super Admin'"), ('Vendor CMS', "'Vendor CMS'"),
              ('End-to-End Workflows', "'End-to-End Workflows'")]
    heads = ['Portal', 'Total cases', 'Executed (Pass + Fail)', 'Passed', 'Failed', 'Blocked', 'Not Run',
             'Not Applicable', 'Pass %', 'Critical cases', 'Smoke cases']
    for i, h in enumerate(heads, 1):
        c = ws.cell(row=4, column=i, value=h)
        c.fill, c.font, c.border = HEAD, HEAD_FONT, BOX
        c.alignment = Alignment(wrap_text=True)
        ws.column_dimensions[get_column_letter(i)].width = 16
    ws.column_dimensions['A'].width = 34
    for r, (label, ref) in enumerate(sheets, 5):
        ws.cell(row=r, column=1, value=label)
        ws.cell(row=r, column=2, value=f'=COUNTA({ref}!A:A)-1')
        ws.cell(row=r, column=3, value=f'=D{r}+E{r}')
        for col, status in zip('DEFGH', ['Pass', 'Fail', 'Blocked', 'Not Run', 'Not Applicable']):
            ws[f'{col}{r}'] = f'=COUNTIF({ref}!L:L,"{status}")'
        ws[f'I{r}'] = f'=IF(C{r}=0,0,D{r}/C{r})'
        ws[f'J{r}'] = f'=COUNTIF({ref}!M:M,"Critical")'
        ws[f'K{r}'] = f'=COUNTIF({ref}!T:T,"Yes")'
    total = 5 + len(sheets)
    ws.cell(row=total, column=1, value='All portals').font = Font(bold=True)
    for col in 'BCDEFGHJK':
        ws[f'{col}{total}'] = f'=SUM({col}5:{col}{total - 1})'
        ws[f'{col}{total}'].font = Font(bold=True)
    ws[f'I{total}'] = f'=IF(C{total}=0,0,D{total}/C{total})'
    for r in range(5, total + 1):
        ws[f'I{r}'].number_format = '0.0%'
        for col in range(1, len(heads) + 1):
            ws.cell(row=r, column=col).border = BOX

    start = total + 3
    ws.cell(row=start - 1, column=1, value='Coverage by module').font = Font(bold=True, size=12)
    mheads = ['Module', 'Portal', 'Cases', 'Passed', 'Failed', 'Blocked', 'Not Run', 'Pass %',
              'Critical', 'Regression script']
    for i, h in enumerate(mheads, 1):
        c = ws.cell(row=start, column=i, value=h)
        c.fill, c.font, c.border = HEAD, HEAD_FONT, BOX
    ws.column_dimensions['J'].width = 48
    sheet_of = {'Super Admin': "'Super Admin'", 'Vendor CMS': "'Vendor CMS'", 'End-to-End': "'End-to-End Workflows'"}
    for r, m in enumerate(MODULES, start + 1):
        ref = sheet_of[m.portal]
        label = module_label(m)
        ws.cell(row=r, column=1, value=label)
        ws.cell(row=r, column=2, value=m.portal)
        ws.cell(row=r, column=3, value=f'=COUNTIF({ref}!C:C,A{r})')
        for col, status in zip('DEFG', ['Pass', 'Fail', 'Blocked', 'Not Run']):
            ws[f'{col}{r}'] = f'=COUNTIFS({ref}!C:C,A{r},{ref}!L:L,"{status}")'
        ws[f'H{r}'] = f'=IF((D{r}+E{r})=0,0,D{r}/(D{r}+E{r}))'
        ws[f'H{r}'].number_format = '0%'
        ws[f'I{r}'] = f'=COUNTIFS({ref}!C:C,A{r},{ref}!M:M,"Critical")'
        ws.cell(row=r, column=10, value=script_path(m))
        for col in range(1, 11):
            ws.cell(row=r, column=col).border = BOX
    last = start + len(MODULES)
    ws.conditional_formatting.add(f'E{start + 1}:E{last}', CellIsRule(
        operator='greaterThan', formula=['0'], fill=PatternFill('solid', fgColor='FFC7CE')))
    ws.conditional_formatting.add(f'H{start + 1}:H{last}', FormulaRule(
        formula=[f'AND(D{start + 1}+E{start + 1}>0,H{start + 1}=1)'], fill=PatternFill('solid', fgColor='C6EFCE')))
    ws.freeze_panes = 'A5'


def defect_sheet(wb) -> None:
    ws = wb.create_sheet('Defect Log')
    cols = [('Defect ID', 12), ('Linked Test Case', 14), ('Portal', 13), ('Module', 26),
            ('Description', 44), ('Severity', 11), ('Steps to Reproduce', 50), ('Expected Result', 40),
            ('Actual Result', 40), ('Status', 13), ('Reported By', 16), ('Date', 12), ('Remarks', 34)]
    header(ws, cols)
    for r, d in enumerate(DEFECTS, 2):
        for col, key in enumerate(['id', 'case', 'portal', 'module', 'description', 'severity', 'steps',
                                   'expected', 'actual', 'status', 'by', 'date', 'remarks'], 1):
            cell = ws.cell(row=r, column=col, value=d.get(key, ''))
            cell.alignment, cell.border = WRAP, BOX
    last = max(len(DEFECTS) + 1, 2)
    ws.auto_filter.ref = f'A1:M{last + 200}'
    list_validation(ws, PRIORITIES, f'F2:F{last + 200}')
    list_validation(ws, ['Open', 'In progress', 'Fixed', 'Retest', 'Closed', "Won't fix"], f'J2:J{last + 200}')
    list_validation(ws, PORTALS, f'C2:C{last + 200}')
    ws.conditional_formatting.add(f'J2:J{last + 200}', CellIsRule(
        operator='equal', formula=['"Open"'], fill=PatternFill('solid', fgColor='FFC7CE')))
    ws.conditional_formatting.add(f'J2:J{last + 200}', CellIsRule(
        operator='equal', formula=['"Closed"'], fill=PatternFill('solid', fgColor='C6EFCE')))


def data_sheet(wb) -> None:
    ws = wb.create_sheet('Test Data & Preconditions')
    cols = [('Item', 30), ('Value', 60), ('Used by', 30), ('Notes', 60)]
    header(ws, cols)
    rows = [
        ('Environment', 'Local QA: admin http://localhost:3000/admin; store http://home-orbit.localhost:3000; database tenantecom_qa (Docker Mongo replica set rs0)', 'All', 'Start the `app-qa` launch entry or `pnpm dev:fresh` with DATABASE_URI pointing at tenantecom_qa. Never run destructive cases on real data.'),
        ('Browser', 'Chrome (or Edge/Firefox); *.localhost opens without setup', 'All', 'Use a normal window for the super admin and private windows or other browsers for store staff and shoppers.'),
        ('Super admin', 'admin@tenantecom.local', 'SA-*, E2E-*', 'Password: SEED_SUPER_ADMIN_PASSWORD in .env (never in this sheet). Two-step required: the tester\'s own authenticator app.'),
        ('Support teammate', 'support.qa@tenantecom.local (invite in SA-14)', 'SA-01, SA-15', 'Sets two-step at first sign-in.'),
        ('Home Orbit owner', 'owner@homeorbit.example', 'VC-*, E2E-*', 'Set the password from the invite link (SA-11 Resend invite prints it in the `pnpm dev` terminal), e.g. HomeOrbit@2026.'),
        ('Demo Sanitary owner', 'owner@demo-sanitary.example', 'E2E-11', 'Invite link from the seed output or Resend invite.'),
        ('Demo Clothing owner', 'owner@demo-clothing.example', 'SA-01 lockout', 'Use for the lockout case so no real account is locked.'),
        ('Store staff to invite', 'editor@ (Catalog editor), orders@ (Order manager), content@ (Content editor), manager@ (Manager) @homeorbit.example', 'VC-28, E2E-10', 'Starter allows 3 staff: remove one before inviting the next, or switch the plan to Enterprise.'),
        ('Shopper', 'Rahul Kulkarni, rahul.k@example.com, mobile 98765 43210', 'E2E-04, E2E-17', 'Login codes print in the `pnpm dev` terminal ("Your code to log in to Home Orbit").'),
        ('Delivery pincodes', '411045 (Pune, Maharashtra), 110001 (Delhi), 744101 (remote zone), 12345 (invalid)', 'VC-24, E2E-04', ''),
        ('GSTINs', 'Valid 29AABCT1234F1ZM, 27AAPFU0939F1ZV; invalid checksum 29AABCT1234F1ZN', 'SA-04, SA-05', 'Sample GSTINs from the GST portal documentation, not real businesses.'),
        ('Products', 'HOPH-504 Feather pull handle; HOGDH-610 Lotus glass door handle; HOAL-101 aldrop', 'VC-04, E2E-03', 'Home Orbit catalogue from `pnpm seed:home-orbit` (110 products).'),
        ('Sample prices and COD', '`pnpm demo:selling home-orbit`', 'Selling cases', 'Local only: puts sample prices, stock, COD (₹49) and two zones on the store.'),
        ('Sample orders', '`pnpm demo:orders home-orbit`', 'SA-02, VC-02', '24 COD orders over 14 days for the dashboards.'),
        ('Razorpay test keys', 'rzp_test_… (your own test keys, optional)', 'VC-25, E2E-09', 'Made-up keys test the refusal path; never live keys.'),
        ('Coupons and schemes', 'HOME200 (₹200 off above ₹1,000); Diwali 2026 (10% off, max ₹1,500)', 'VC-14, VC-15, E2E-05', 'Created during the cases.'),
        ('Emails and WhatsApp', 'Printed in the `pnpm dev` terminal as [dev-log email] / [dev-log whatsapp]', 'Messaging cases', 'Nothing is really sent locally.'),
        ('Reset', 'docker compose down -v; up -d mongo; pnpm seed; pnpm seed:home-orbit (with DATABASE_URI set to the QA database)', 'After destructive runs', 'Uploaded files stay in media/.'),
        ('Run order', 'SA-01 → SA-03 → SA-05/06 → VC-01 → VC-03 → VC-04…; selling cases after `pnpm demo:selling`; E2E after the modules they depend on', 'All', 'See docs/qa/coverage.md "Dependencies".'),
    ]
    for r, row in enumerate(rows, 2):
        for col, v in enumerate(row, 1):
            cell = ws.cell(row=r, column=col, value=v)
            cell.alignment, cell.border = WRAP, BOX


def write_workbook() -> Path:
    wb = Workbook()
    summary_sheet(wb)
    case_sheet(wb, 'Vendor CMS', 'Vendor CMS')
    case_sheet(wb, 'Super Admin', 'Super Admin')
    case_sheet(wb, 'End-to-End Workflows', 'End-to-End')
    defect_sheet(wb)
    data_sheet(wb)
    out = ROOT / 'TenantEcom-test-cases.xlsx'
    wb.save(out)
    return out


if __name__ == '__main__':
    write_scripts()
    write_coverage()
    out = write_workbook()
    per = {p: sum(1 for c in CASES if BY_ID[c.module].portal == p) for p in PORTALS}
    print(f'{len(MODULES)} modules, {len(CASES)} cases {per}; '
          f'{sum(1 for c in CASES if c.smoke)} smoke; {len(RESULTS)} results; {len(DEFECTS)} defects -> {out}')
